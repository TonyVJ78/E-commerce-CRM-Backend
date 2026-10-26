"""
Exportación de los informes con IA (CU-34) a PDF y Excel.

Ambos formatos salen de lo guardado en `ReporteGenerado` (tabla + análisis +
gráfico sugerido), así que exportar no vuelve a consultar la base ni a llamar
a la IA: el archivo coincide con lo que el usuario vio en pantalla.

- PDF con fpdf2 (Python puro, corre en Vercel; mismo enfoque que el recibo de CU-33).
- Excel con openpyxl: hoja "Datos" con formato numérico, filtros y un gráfico
  nativo de Excel, y hoja "Análisis" con el texto.
"""

import io
import re

from django.utils import timezone
from django.utils.text import slugify
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.pedidos.recibo import _t

from .informes_consultas import DECIMAL, ENTERO, MONEDA


ROJO = (200, 16, 46)        # --primary #C8102E
AMARILLO = (244, 208, 63)   # --accent  #F4D03F
GRIS_TEXTO = (75, 85, 99)
GRIS_LINEA = (229, 231, 235)
GRIS_FONDO = (249, 250, 251)
MAX_BARRAS_PDF = 12

NUMERICOS = (ENTERO, MONEDA, DECIMAL)


def nombre_archivo(informe, extension):
    base = slugify(informe.titulo or 'informe')[:60] or 'informe'
    fecha = timezone.localtime(informe.fecha_generacion).strftime('%Y%m%d')
    return f'informe-{base}-{fecha}.{extension}'


def _alcance(informe):
    return informe.tienda.nombre if informe.tienda else 'Varias tiendas'


def _formato(valor, tipo):
    if valor is None:
        return '-'
    if tipo == MONEDA:
        return f'Bs {valor:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    if tipo == ENTERO:
        return f'{int(valor):,}'.replace(',', '.')
    if tipo == DECIMAL:
        return f'{valor:.2f}'.replace('.', ',')
    return str(valor)


# ---------------------------------------------------------------------------
# Markdown mínimo → líneas para el PDF
# ---------------------------------------------------------------------------

_NEGRITA = re.compile(r'\*\*(.+?)\*\*')
_ENLACE = re.compile(r'\[([^\]]+)\]\([^)]+\)')
_RESTO_MD = re.compile(r'(?<!\w)[*_`]+|[*_`]+(?!\w)')


def _lineas_markdown(texto):
    """Convierte el Markdown del análisis en `(tipo, texto)` con negritas `**`.

    `tipo` es `titulo`, `vineta` o `parrafo`. Las negritas se conservan como
    `**...**` porque `multi_cell(markdown=True)` de fpdf2 las entiende.
    """
    resultado = []
    for cruda in (texto or '').splitlines():
        linea = _ENLACE.sub(r'\1', cruda.strip())
        if not linea:
            resultado.append(('vacia', ''))
            continue
        tipo = 'parrafo'
        if linea.startswith('#'):
            tipo, linea = 'titulo', linea.lstrip('#').strip()
        elif re.match(r'^([-*+]|\d+[.)])\s+', linea):
            numero = re.match(r'^(\d+[.)])\s+', linea)
            linea = re.sub(r'^([-*+]|\d+[.)])\s+', '', linea)
            linea = f'{numero.group(1)} {linea}' if numero else f'- {linea}'
            tipo = 'vineta'
        # Negritas: se protegen, se limpia el resto de marcas y se restauran.
        partes = _NEGRITA.split(linea)
        limpias = [(_RESTO_MD.sub('', p) if i % 2 == 0 else f'**{p}**') for i, p in enumerate(partes)]
        resultado.append((tipo, ''.join(limpias)))
    return resultado


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

class _InformePDF(FPDF):
    titulo_pie = ''

    def footer(self):
        self.set_y(-14)
        self.set_draw_color(*GRIS_LINEA)
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(1.5)
        self.set_font('Helvetica', '', 7.5)
        self.set_text_color(*GRIS_TEXTO)
        self.cell(0, 4, _t(f'Kantu Market · {self.titulo_pie}'), align='L')
        self.set_x(self.l_margin)
        self.cell(0, 4, _t(f'Página {self.page_no()}/{{nb}}'), align='R')


def _anchos_columnas(pdf, columnas, filas, ancho_total):
    pdf.set_font('Helvetica', '', 8)
    deseados = []
    for i, col in enumerate(columnas):
        muestras = [col['titulo']] + [_formato(f[i], col['tipo']) for f in filas[:80]]
        deseados.append(min(max(pdf.get_string_width(_t(m)) for m in muestras) + 5, 80))
    factor = ancho_total / sum(deseados)
    return [d * factor for d in deseados]


def _recortar(pdf, texto, ancho):
    texto = _t(texto)
    if pdf.get_string_width(texto) <= ancho:
        return texto
    while texto and pdf.get_string_width(texto + '...') > ancho:
        texto = texto[:-1]
    return texto + '...'


def _encabezado_tabla(pdf, columnas, anchos):
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_fill_color(*ROJO)
    pdf.set_text_color(255, 255, 255)
    for col, ancho in zip(columnas, anchos):
        alinear = 'R' if col['tipo'] in NUMERICOS else 'L'
        pdf.cell(ancho, 7, _recortar(pdf, col['titulo'], ancho - 2), fill=True, align=alinear)
    pdf.ln()
    pdf.set_text_color(17, 24, 39)
    pdf.set_font('Helvetica', '', 8)


def _tabla(pdf, columnas, filas):
    ancho_total = pdf.w - pdf.l_margin - pdf.r_margin
    anchos = _anchos_columnas(pdf, columnas, filas, ancho_total)
    _encabezado_tabla(pdf, columnas, anchos)
    for n, fila in enumerate(filas):
        if pdf.will_page_break(6.5):
            pdf.add_page()
            _encabezado_tabla(pdf, columnas, anchos)
        pdf.set_fill_color(*(GRIS_FONDO if n % 2 else (255, 255, 255)))
        for col, ancho, valor in zip(columnas, anchos, fila):
            alinear = 'R' if col['tipo'] in NUMERICOS else 'L'
            pdf.cell(ancho, 6.5, _recortar(pdf, _formato(valor, col['tipo']), ancho - 2), fill=True, align=alinear)
        pdf.ln()

    # Totales de las columnas que suman (unidades, montos)
    sumables = [i for i, c in enumerate(columnas) if c.get('sumable')]
    if filas and sumables and len(filas) > 1:
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_draw_color(*GRIS_TEXTO)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + ancho_total, pdf.get_y())
        for i, (col, ancho) in enumerate(zip(columnas, anchos)):
            if i in sumables:
                total = sum(f[i] or 0 for f in filas)
                pdf.cell(ancho, 6.5, _recortar(pdf, _formato(total, col['tipo']), ancho - 2), align='R')
            else:
                pdf.cell(ancho, 6.5, 'Total' if i == 0 else '', align='L')
        pdf.ln()


def _grafico_pdf(pdf, informe, columnas, filas):
    grafico = informe.grafico or {}
    claves = [c['clave'] for c in columnas]
    if grafico.get('tipo') not in ('barras', 'lineas') or grafico.get('etiqueta') not in claves \
            or grafico.get('valor') not in claves or not filas:
        return
    i_etq, i_val = claves.index(grafico['etiqueta']), claves.index(grafico['valor'])
    col_val = columnas[i_val]
    datos = [(str(f[i_etq]), float(f[i_val] or 0)) for f in filas]
    # En series de tiempo se muestran los últimos periodos; en rankings, los primeros.
    datos = datos[-MAX_BARRAS_PDF:] if grafico['tipo'] == 'lineas' else datos[:MAX_BARRAS_PDF]
    maximo = max((v for _, v in datos), default=0) or 1

    ancho = pdf.w - pdf.l_margin - pdf.r_margin
    alto_barra = 5.5
    if pdf.will_page_break(len(datos) * (alto_barra + 1.5) + 14):
        pdf.add_page()
    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(17, 24, 39)
    pdf.cell(0, 7, _t(f'{col_val["titulo"]} por {columnas[i_etq]["titulo"].lower()}'),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1)
    ancho_etq, ancho_valor = ancho * 0.30, ancho * 0.18
    ancho_barras = ancho - ancho_etq - ancho_valor
    pdf.set_font('Helvetica', '', 8)
    for etiqueta, valor in datos:
        y = pdf.get_y()
        pdf.set_text_color(*GRIS_TEXTO)
        pdf.cell(ancho_etq, alto_barra, _recortar(pdf, etiqueta, ancho_etq - 2))
        largo = max(ancho_barras * (valor / maximo), 0.6) if valor > 0 else 0
        pdf.set_fill_color(*ROJO)
        if largo:
            pdf.rect(pdf.l_margin + ancho_etq, y + 0.8, largo, alto_barra - 1.6, style='F')
        pdf.set_xy(pdf.l_margin + ancho_etq + ancho_barras, y)
        pdf.set_text_color(17, 24, 39)
        pdf.cell(ancho_valor, alto_barra, _t(_formato(valor, col_val['tipo'])), align='R',
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(1.5)
    pdf.ln(3)


def generar_pdf(informe):
    datos = informe.datos or {}
    columnas, filas = datos.get('columnas', []), datos.get('filas', [])
    # Apaisado si la tabla es ancha
    orientacion = 'L' if len(columnas) > 5 else 'P'

    pdf = _InformePDF(orientation=orientacion, unit='mm', format='A4')
    pdf.titulo_pie = informe.titulo
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(14, 14, 14)
    pdf.alias_nb_pages()
    pdf.add_page()

    # Banda de marca
    pdf.set_fill_color(*ROJO)
    pdf.rect(0, 0, pdf.w, 26, style='F')
    pdf.set_fill_color(*AMARILLO)
    pdf.rect(0, 26, pdf.w, 1.4, style='F')
    pdf.set_xy(14, 7)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 16)
    pdf.cell(0, 7, _t('Kantu Market'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 9)
    pdf.cell(0, 5, _t('Informe generado con IA'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_y(33)
    pdf.set_text_color(17, 24, 39)
    pdf.set_font('Helvetica', 'B', 14)
    pdf.multi_cell(0, 7, _t(informe.titulo), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font('Helvetica', '', 8.5)
    pdf.set_text_color(*GRIS_TEXTO)
    fecha = timezone.localtime(informe.fecha_generacion).strftime('%d/%m/%Y %H:%M')
    meta = [f'Alcance: {_alcance(informe)}', f'Generado: {fecha}',
            f'Por: {informe.generado_por_usuario.email}']
    pdf.multi_cell(0, 4.5, _t(' · '.join(meta)), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if datos.get('filtros'):
        pdf.multi_cell(0, 4.5, _t(datos['filtros']), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if informe.pregunta:
        pdf.set_font('Helvetica', 'I', 8.5)
        pdf.multi_cell(0, 4.5, _t(f'Pedido: "{informe.pregunta}"'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    # Análisis
    if informe.analisis:
        pdf.set_text_color(*ROJO)
        pdf.set_font('Helvetica', 'B', 11)
        pdf.cell(0, 7, _t('Análisis'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_text_color(17, 24, 39)
        for tipo, linea in _lineas_markdown(informe.analisis):
            if tipo == 'vacia':
                pdf.ln(1.5)
            elif tipo == 'titulo':
                pdf.set_font('Helvetica', 'B', 10)
                pdf.multi_cell(0, 5.5, _t(linea.replace('**', '')), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            else:
                pdf.set_font('Helvetica', '', 9.5)
                sangria = 4 if tipo == 'vineta' else 0
                pdf.set_x(pdf.l_margin + sangria)
                pdf.multi_cell(0, 5, _t(linea), markdown=True, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(4)

    _grafico_pdf(pdf, informe, columnas, filas)

    # Tabla
    pdf.set_text_color(*ROJO)
    pdf.set_font('Helvetica', 'B', 11)
    pdf.cell(0, 7, _t('Datos'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if columnas and filas:
        _tabla(pdf, columnas, filas)
    else:
        pdf.set_font('Helvetica', 'I', 9)
        pdf.set_text_color(*GRIS_TEXTO)
        pdf.cell(0, 6, _t('La consulta no devolvió filas.'), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if datos.get('truncado'):
        pdf.ln(2)
        pdf.set_font('Helvetica', 'I', 8)
        pdf.set_text_color(*GRIS_TEXTO)
        pdf.multi_cell(0, 4.5, _t(
            f'Se muestran {len(filas)} de {datos.get("total_filas")} filas que cumplen los filtros.'))

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_FORMATO_EXCEL = {
    MONEDA: '"Bs" #,##0.00',
    ENTERO: '#,##0',
    DECIMAL: '0.00',
}


def generar_xlsx(informe):
    datos = informe.datos or {}
    columnas, filas = datos.get('columnas', []), datos.get('filas', [])

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Datos'

    rojo = PatternFill('solid', fgColor='C8102E')
    blanco_negrita = Font(bold=True, color='FFFFFF')

    hoja['A1'] = informe.titulo
    hoja['A1'].font = Font(bold=True, size=14, color='C8102E')
    fecha = timezone.localtime(informe.fecha_generacion).strftime('%d/%m/%Y %H:%M')
    hoja['A2'] = f'Alcance: {_alcance(informe)} · Generado: {fecha} · Por: {informe.generado_por_usuario.email}'
    hoja['A3'] = datos.get('filtros') or ''
    for celda in ('A2', 'A3'):
        hoja[celda].font = Font(italic=True, color='4B5563', size=9)

    fila_encabezado = 5
    for j, col in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_encabezado, column=j, value=col['titulo'])
        celda.fill, celda.font = rojo, blanco_negrita
        celda.alignment = Alignment(horizontal='right' if col['tipo'] in NUMERICOS else 'left', vertical='center')

    for i, fila in enumerate(filas, start=fila_encabezado + 1):
        for j, (col, valor) in enumerate(zip(columnas, fila), start=1):
            celda = hoja.cell(row=i, column=j, value=valor)
            if col['tipo'] in _FORMATO_EXCEL:
                celda.number_format = _FORMATO_EXCEL[col['tipo']]

    ultima = fila_encabezado + len(filas)
    sumables = [j for j, c in enumerate(columnas, start=1) if c.get('sumable')]
    if len(filas) > 1 and sumables:
        fila_total = ultima + 1
        hoja.cell(row=fila_total, column=1, value='Total').font = Font(bold=True)
        for j in sumables:
            letra = get_column_letter(j)
            celda = hoja.cell(row=fila_total, column=j,
                              value=f'=SUBTOTAL(109,{letra}{fila_encabezado + 1}:{letra}{ultima})')
            celda.font = Font(bold=True)
            celda.number_format = _FORMATO_EXCEL[columnas[j - 1]['tipo']]

    if columnas:
        hoja.freeze_panes = hoja.cell(row=fila_encabezado + 1, column=1)
        if filas:
            hoja.auto_filter.ref = f'A{fila_encabezado}:{get_column_letter(len(columnas))}{ultima}'
    for j, col in enumerate(columnas, start=1):
        largo = max([len(str(col['titulo']))] + [len(str(f[j - 1] or '')) for f in filas[:200]])
        hoja.column_dimensions[get_column_letter(j)].width = min(max(largo + 3, 10), 50)

    if datos.get('truncado'):
        hoja.cell(row=ultima + 3, column=1,
                  value=f'Se muestran {len(filas)} de {datos.get("total_filas")} filas que cumplen los filtros.'
                  ).font = Font(italic=True, color='4B5563', size=9)

    # Gráfico nativo de Excel
    grafico = informe.grafico or {}
    claves = [c['clave'] for c in columnas]
    if filas and grafico.get('tipo') in ('barras', 'lineas') \
            and grafico.get('etiqueta') in claves and grafico.get('valor') in claves:
        j_etq, j_val = claves.index(grafico['etiqueta']) + 1, claves.index(grafico['valor']) + 1
        hasta = min(ultima, fila_encabezado + 30)
        chart = LineChart() if grafico['tipo'] == 'lineas' else BarChart()
        if grafico['tipo'] == 'barras':
            chart.type = 'bar' if len(filas) > 8 else 'col'
        chart.title = f'{columnas[j_val - 1]["titulo"]} por {columnas[j_etq - 1]["titulo"].lower()}'
        chart.legend = None
        chart.height, chart.width = 9, 18
        chart.add_data(Reference(hoja, min_col=j_val, min_row=fila_encabezado, max_row=hasta), titles_from_data=True)
        chart.set_categories(Reference(hoja, min_col=j_etq, min_row=fila_encabezado + 1, max_row=hasta))
        hoja.add_chart(chart, f'{get_column_letter(len(columnas) + 2)}{fila_encabezado}')

    # Hoja de análisis
    analisis = libro.create_sheet('Análisis')
    analisis.column_dimensions['A'].width = 110
    analisis['A1'] = informe.titulo
    analisis['A1'].font = Font(bold=True, size=14, color='C8102E')
    fila = 3
    if informe.pregunta:
        analisis.cell(row=fila, column=1, value=f'Pedido: "{informe.pregunta}"').font = Font(italic=True, color='4B5563')
        fila += 2
    for tipo, linea in _lineas_markdown(informe.analisis):
        celda = analisis.cell(row=fila, column=1, value=linea.replace('**', ''))
        celda.alignment = Alignment(wrap_text=True, vertical='top')
        if tipo == 'titulo':
            celda.font = Font(bold=True, size=12)
        fila += 1

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
