"""Pruebas de los informes de tienda con IA (CU-34).

La API de Claude se simula con `mock`: no se hace ninguna llamada de red. Los
pedidos se crean como en el checkout real: el trigger de la base recalcula el
total y descuenta el stock.
"""

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest import mock

from django.core.cache import cache
from django.test import override_settings
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.catalogo.models import Categoria, Producto, Variante
from apps.pedidos.models import ItemPedido, Pedido
from apps.tiendas.models import Tienda, UsuarioTienda
from apps.usuarios.models import Permiso, Rol, RolPermiso, Usuario

from . import informes_consultas as consultas
from .informes_exportar import _lineas_markdown, generar_pdf, generar_xlsx
from .models import ReporteGenerado


URL_GENERAR = '/api/ia/informes/generar/'
URL_CONTEXTO = '/api/ia/informes/contexto/'
URL_LISTA = '/api/ia/informes/'


def texto(contenido):
    return SimpleNamespace(type='text', text=contenido)


def uso_herramienta(nombre, entrada, id_='toolu_1'):
    return SimpleNamespace(type='tool_use', name=nombre, input=entrada, id=id_)


def respuesta(*bloques, stop_reason='end_turn'):
    return SimpleNamespace(content=list(bloques), stop_reason=stop_reason)


def mensajes(*textos):
    return {'mensajes': [{'rol': 'usuario', 'contenido': t} for t in textos]}


class BaseInformes(APITestCase):
    @classmethod
    def setUpTestData(cls):
        rol_empresa = Rol.objects.get_or_create(nombre='empresa')[0]
        rol_admin = Rol.objects.get_or_create(nombre='administrador')[0]
        rol_cliente = Rol.objects.get_or_create(nombre='cliente')[0]

        cls.empresa = Usuario.objects.create_user(
            email='empresa-inf@example.com', password='Password123!', rol=rol_empresa)
        cls.otra_empresa = Usuario.objects.create_user(
            email='otra-inf@example.com', password='Password123!', rol=rol_empresa)
        cls.admin = Usuario.objects.create_user(
            email='admin-inf@example.com', password='Password123!', rol=rol_admin)
        cls.cliente = Usuario.objects.create_user(
            email='cliente-inf@example.com', password='Password123!', rol=rol_cliente)
        # Encargado: rol cliente (sin permiso de informes) asignado a una tienda.
        cls.encargado = Usuario.objects.create_user(
            email='encargado-inf@example.com', password='Password123!', rol=rol_cliente)

        cls.tienda = Tienda.objects.create(propietario=cls.empresa, nombre='Alpaca Andina')
        cls.tienda_ajena = Tienda.objects.create(propietario=cls.otra_empresa, nombre='Tienda Ajena')
        UsuarioTienda.objects.create(usuario=cls.encargado, tienda=cls.tienda, rol_interno='encargado')

        textiles = Categoria.objects.create(tienda=cls.tienda, nombre='Textiles')
        cls.chompa = cls._variante(cls.tienda, textiles, 'Chompa de Alpaca', Decimal('350.00'), stock=50)
        cls.gorro = cls._variante(cls.tienda, textiles, 'Gorro andino', Decimal('80.00'), stock=3,
                                  stock_minimo=5)
        cls.ajena = cls._variante(cls.tienda_ajena, None, 'Producto ajeno', Decimal('999.00'), stock=50)

        cls._pedido(cls.tienda, [(cls.chompa, 2), (cls.gorro, 1)])        # 780
        cls._pedido(cls.tienda, [(cls.chompa, 1)])                        # 350
        cls._pedido(cls.tienda, [(cls.gorro, 1)], estado='cancelado')     # no cuenta
        cls._pedido(cls.tienda_ajena, [(cls.ajena, 5)])                   # otra tienda

    @classmethod
    def _variante(cls, tienda, categoria, nombre, precio, stock, stock_minimo=5):
        producto = Producto.objects.create(tienda=tienda, categoria=categoria, nombre=nombre, slug=nombre.lower())
        return Variante.objects.create(producto=producto, sku=f'SKU-{producto.id}', precio=precio,
                                       stock=stock, stock_minimo=stock_minimo)

    @classmethod
    def _pedido(cls, tienda, lineas, estado='pendiente'):
        pedido = Pedido.objects.create(cliente=cls.cliente, tienda=tienda, subtotal=0, total=0,
                                       estado_actual='pendiente')
        for variante, cantidad in lineas:
            ItemPedido.objects.create(tienda=tienda, pedido=pedido, variante=variante,
                                      cantidad=cantidad, precio_unitario=variante.precio)
        if estado != 'pendiente':
            Pedido.objects.filter(pk=pedido.pk).update(estado_actual=estado)
        return pedido


class MotorConsultasTests(BaseInformes):
    def _ejecutar(self, spec, tiendas=None):
        tiendas = tiendas or [self.tienda.id]
        return consultas.ejecutar(consultas.normalizar(spec), tiendas)

    def test_ventas_por_producto_excluye_cancelados_y_otras_tiendas(self):
        r = self._ejecutar({'fuente': 'ventas_productos', 'agrupar_por': ['producto'],
                            'metricas': ['unidades', 'ingresos']})
        self.assertEqual([c['clave'] for c in r['columnas']], ['producto', 'unidades', 'ingresos'])
        self.assertEqual(r['filas'], [['Chompa de Alpaca', 3, 1050.0], ['Gorro andino', 1, 80.0]])

    def test_conteos_distintos_y_promedios_no_son_sumables(self):
        r = self._ejecutar({'fuente': 'ventas_productos', 'agrupar_por': ['producto'],
                            'metricas': ['unidades', 'pedidos', 'precio_promedio']})
        self.assertEqual([c['sumable'] for c in r['columnas']], [False, True, False, False])
        self.assertEqual(consultas.resumen_para_modelo(r)['totales_de_la_tabla'], {'unidades': 4})

    def test_totales_sin_agrupar(self):
        r = self._ejecutar({'fuente': 'pedidos', 'metricas': ['pedidos', 'monto_total', 'ticket_promedio']})
        self.assertEqual(r['filas'], [[2, 1130.0, 565.0]])

    def test_filtro_de_estado_incluye_cancelados(self):
        r = self._ejecutar({'fuente': 'pedidos', 'agrupar_por': ['estado'], 'metricas': ['pedidos'],
                            'filtros': {'estados': ['cancelado']}})
        self.assertEqual(r['filas'], [['Cancelado', 1]])

    def test_inventario_stock_bajo(self):
        r = self._ejecutar({'fuente': 'inventario', 'agrupar_por': ['producto'], 'metricas': ['stock'],
                            'filtros': {'solo_stock_bajo': True}})
        self.assertEqual([f[0] for f in r['filas']], ['Gorro andino'])

    def test_serie_por_mes_ordena_por_tiempo(self):
        r = self._ejecutar({'fuente': 'pedidos', 'agrupar_por': ['mes'], 'metricas': ['monto_total']})
        self.assertEqual(len(r['filas']), 1)
        self.assertRegex(r['filas'][0][0], r'^\d{4}-\d{2}$')

    def test_limite_marca_truncado(self):
        r = self._ejecutar({'fuente': 'ventas_productos', 'agrupar_por': ['producto'],
                            'metricas': ['unidades'], 'limite': 1})
        self.assertEqual(len(r['filas']), 1)
        self.assertTrue(r['truncado'])
        self.assertEqual(r['total_filas'], 2)

    def test_filtros_fuera_de_su_objeto_se_aplican(self):
        # El modelo a veces pone fecha_desde junto a `fuente`: debe filtrar igual.
        spec = consultas.normalizar({'fuente': 'pedidos', 'metricas': ['pedidos'], 'fecha_desde': '2999-01-01'})
        self.assertEqual(spec['filtros'], {'fecha_desde': '2999-01-01'})
        self.assertEqual(consultas.ejecutar(spec, [self.tienda.id])['filas'], [[0]])
        self.assertIn('Periodo: 2999-01-01', consultas.describir_filtros(spec, {}))

    def test_sin_fecha_el_periodo_es_todo_el_historial(self):
        spec = consultas.normalizar({'fuente': 'pedidos', 'metricas': ['pedidos']})
        self.assertIn('Periodo: todo el historial', consultas.describir_filtros(spec, {}))
        resumen = consultas.resumen_para_modelo(consultas.ejecutar(spec, [self.tienda.id]), spec)
        self.assertIn('todo el historial', resumen['filtros_aplicados'])

    def test_especificaciones_invalidas(self):
        casos = [
            {'fuente': 'usuarios', 'metricas': ['x']},
            {'fuente': 'pedidos', 'metricas': ['stock']},
            {'fuente': 'pedidos', 'agrupar_por': ['sku'], 'metricas': ['pedidos']},
            {'fuente': 'inventario', 'metricas': ['stock'], 'filtros': {'fecha_desde': '2026-01-01'}},
            {'fuente': 'pedidos', 'metricas': ['pedidos'], 'filtros': {'sql': 'drop table'}},
            {'fuente': 'pedidos', 'metricas': ['pedidos'], 'sql': 'drop table'},
            {'fuente': 'pedidos', 'metricas': ['pedidos'], 'filtros': {'fecha_desde': 'ayer'}},
        ]
        for spec in casos:
            with self.subTest(spec=spec), self.assertRaises(consultas.ConsultaInvalida):
                self._ejecutar(spec)

    def test_alcance_por_rol(self):
        self.assertEqual(set(consultas.tiendas_permitidas(self.empresa)), {self.tienda})
        self.assertEqual(set(consultas.tiendas_permitidas(self.encargado)), {self.tienda})
        self.assertEqual(set(consultas.tiendas_permitidas(self.admin)), {self.tienda, self.tienda_ajena})


@override_settings(ANTHROPIC_API_KEY='sk-ant-test', INFORMES_MODEL='claude-haiku-4-5')
class InformesApiTests(BaseInformes):
    def setUp(self):
        cache.clear()  # el throttle guarda sus contadores en la caché
        parche = mock.patch('apps.ia.informes.anthropic.Anthropic')
        self.Anthropic = parche.start()
        self.addCleanup(parche.stop)
        self.crear = self.Anthropic.return_value.messages.create

    def _consulta_y_presenta(self, spec, **presentar):
        self.crear.side_effect = [
            respuesta(uso_herramienta('consultar_datos', spec, 'toolu_1'), stop_reason='tool_use'),
            respuesta(uso_herramienta('presentar_informe', {
                'consulta_id': 1, 'titulo': 'Más vendidos', 'analisis': '- La **chompa** lidera.',
                'grafico': {'tipo': 'barras', 'etiqueta': 'producto', 'valor': 'unidades'},
                'sugerencias': ['Ventas por semana'], **presentar,
            }, 'toolu_2'), stop_reason='tool_use'),
        ]

    def test_cliente_sin_permiso_recibe_403(self):
        self.client.force_authenticate(self.cliente)
        self.assertEqual(self.client.get(URL_CONTEXTO).status_code, 403)
        self.assertEqual(self.client.post(URL_GENERAR, mensajes('hola'), format='json').status_code, 403)

    def test_permiso_se_concede_desde_la_matriz(self):
        # El encargado tiene rol cliente: con el permiso asignado al rol ya puede.
        rol_cliente = Rol.objects.get(nombre='cliente')
        for codigo in ('informes.ver', 'informes.crear'):
            RolPermiso.objects.create(rol=rol_cliente, permiso=Permiso.objects.get(codigo=codigo))
        self.client.force_authenticate(self.encargado)
        res = self.client.get(URL_CONTEXTO)
        self.assertEqual(res.status_code, 200)
        self.assertEqual([t['id'] for t in res.data['tiendas']], [self.tienda.id])

    def test_contexto_de_empresa(self):
        self.client.force_authenticate(self.empresa)
        res = self.client.get(URL_CONTEXTO)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['tiendas'], [{'id': self.tienda.id, 'nombre': 'Alpaca Andina'}])
        self.assertTrue(res.data['puede_crear'])
        self.assertFalse(res.data['es_administrador'])

    def test_pedido_general_devuelve_preguntas(self):
        self.client.force_authenticate(self.empresa)
        self.crear.return_value = respuesta(texto(
            '¿Qué quieres medir y de qué periodo?\nSUGERENCIAS: Ventas del mes | Stock bajo'))
        res = self.client.post(URL_GENERAR, mensajes('quiero un informe'), format='json')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['tipo'], 'pregunta')
        self.assertEqual(res.data['sugerencias'], ['Ventas del mes', 'Stock bajo'])
        self.assertIsNone(res.data['informe'])
        self.assertFalse(ReporteGenerado.objects.exists())

    def test_genera_informe_con_tabla_de_la_base(self):
        self.client.force_authenticate(self.empresa)
        self._consulta_y_presenta({'fuente': 'ventas_productos', 'agrupar_por': ['producto'],
                                   'metricas': ['unidades']})
        res = self.client.post(URL_GENERAR, mensajes('productos más vendidos'), format='json')

        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['tipo'], 'informe')
        informe = res.data['informe']
        self.assertEqual(informe['titulo'], 'Más vendidos')
        self.assertEqual(informe['filas'], [['Chompa de Alpaca', 3], ['Gorro andino', 1]])
        self.assertEqual(informe['grafico'], {'tipo': 'barras', 'etiqueta': 'producto', 'valor': 'unidades'})
        self.assertEqual(informe['tienda'], self.tienda.id)
        self.assertEqual(res.data['sugerencias'], ['Ventas por semana'])

        guardado = ReporteGenerado.objects.get()
        self.assertEqual(guardado.generado_por_usuario, self.empresa)
        self.assertEqual(guardado.pregunta, 'productos más vendidos')

        # El modelo recibió el resultado de la consulta con su id.
        segunda_llamada = self.crear.call_args_list[1].kwargs['messages']
        self.assertIn('"consulta_id": 1', segunda_llamada[-1]['content'][0]['content'])

    def test_el_modelo_no_puede_salir_de_las_tiendas_del_usuario(self):
        self.client.force_authenticate(self.empresa)
        self.crear.side_effect = [
            respuesta(uso_herramienta('consultar_datos', {
                'fuente': 'pedidos', 'metricas': ['pedidos'],
                'filtros': {'tienda_ids': [self.tienda_ajena.id]}}), stop_reason='tool_use'),
            respuesta(texto('No tengo acceso a esa tienda.\nSUGERENCIAS: Mis ventas')),
        ]
        res = self.client.post(URL_GENERAR, mensajes('ventas de Tienda Ajena'), format='json')
        self.assertEqual(res.data['tipo'], 'pregunta')
        resultado = self.crear.call_args_list[1].kwargs['messages'][-1]['content'][0]
        self.assertTrue(resultado['is_error'])

    def test_tienda_ajena_en_la_peticion_es_403(self):
        self.client.force_authenticate(self.empresa)
        datos = {**mensajes('ventas'), 'tienda_id': self.tienda_ajena.id}
        self.assertEqual(self.client.post(URL_GENERAR, datos, format='json').status_code, 403)
        self.crear.assert_not_called()

    def test_admin_ve_todas_las_tiendas(self):
        self.client.force_authenticate(self.admin)
        self._consulta_y_presenta({'fuente': 'pedidos', 'agrupar_por': ['tienda'], 'metricas': ['pedidos']})
        res = self.client.post(URL_GENERAR, mensajes('pedidos por tienda'), format='json')
        self.assertEqual(sorted(f[0] for f in res.data['informe']['filas']), ['Alpaca Andina', 'Tienda Ajena'])
        self.assertIsNone(res.data['informe']['tienda'])

    def test_si_contesta_en_texto_tras_consultar_igual_arma_el_informe(self):
        self.client.force_authenticate(self.empresa)
        self.crear.side_effect = [
            respuesta(uso_herramienta('consultar_datos', {'fuente': 'pedidos', 'metricas': ['pedidos']}),
                      stop_reason='tool_use'),
            respuesta(texto('Tuviste **2 pedidos**.')),
        ]
        res = self.client.post(URL_GENERAR, mensajes('cuántos pedidos'), format='json')
        self.assertEqual(res.data['tipo'], 'informe')
        self.assertEqual(res.data['informe']['filas'], [[2]])

    def test_sin_clave_responde_503(self):
        self.client.force_authenticate(self.empresa)
        with override_settings(ANTHROPIC_API_KEY=''):
            res = self.client.post(URL_GENERAR, mensajes('ventas'), format='json')
        self.assertEqual(res.status_code, 503)

    def test_historial_exportacion_y_borrado_solo_del_dueno(self):
        self.client.force_authenticate(self.empresa)
        self._consulta_y_presenta({'fuente': 'ventas_productos', 'agrupar_por': ['producto'],
                                   'metricas': ['unidades', 'ingresos']})
        informe_id = self.client.post(URL_GENERAR, mensajes('más vendidos'), format='json').data['informe']['id']

        lista = self.client.get(URL_LISTA)
        self.assertEqual([i['id'] for i in lista.data], [informe_id])

        pdf = self.client.get(f'{URL_LISTA}{informe_id}/exportar/?formato=pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))

        xlsx = self.client.get(f'{URL_LISTA}{informe_id}/exportar/?formato=xlsx')
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn('attachment', xlsx['Content-Disposition'])

        self.client.force_authenticate(self.otra_empresa)
        self.assertEqual(self.client.get(f'{URL_LISTA}{informe_id}/').status_code, 404)
        self.assertEqual(self.client.get(f'{URL_LISTA}{informe_id}/exportar/').status_code, 404)

        self.client.force_authenticate(self.empresa)
        self.assertEqual(self.client.delete(f'{URL_LISTA}{informe_id}/').status_code, 204)
        self.assertFalse(ReporteGenerado.objects.exists())


class ExportarTests(BaseInformes):
    def _informe(self, **extra):
        datos = {
            'columnas': [
                {'clave': 'producto', 'titulo': 'Producto', 'tipo': 'texto', 'sumable': False},
                {'clave': 'unidades', 'titulo': 'Unidades vendidas', 'tipo': 'entero', 'sumable': True},
                {'clave': 'ingresos', 'titulo': 'Ingresos (Bs)', 'tipo': 'moneda', 'sumable': True},
                {'clave': 'pedidos', 'titulo': 'Pedidos', 'tipo': 'entero', 'sumable': False},
            ],
            'filas': [['Chompa “premium” de Alpaca 🦙', 3, 1050.0, 2], ['Gorro andino', 1, 80.0, 1]],
            'total_filas': 2, 'truncado': False, 'filtros': 'Periodo: 2026-09-01 a 2026-09-30',
        }
        return ReporteGenerado.objects.create(
            tienda=self.tienda, generado_por_usuario=self.empresa, tipo_reporte='informe_ia', formato='tabla',
            titulo='Más vendidos de septiembre', pregunta='más vendidos', datos=datos,
            analisis='### Hallazgos\n- La **chompa** lidera con 3 unidades.\n1. Repón stock.',
            grafico={'tipo': 'barras', 'etiqueta': 'producto', 'valor': 'ingresos'}, **extra)

    def test_pdf_con_caracteres_fuera_de_latin1(self):
        self.assertTrue(generar_pdf(self._informe()).startswith(b'%PDF'))

    def test_xlsx_tiene_datos_total_grafico_y_analisis(self):
        libro = load_workbook(BytesIO(generar_xlsx(self._informe())))
        hoja = libro['Datos']
        self.assertEqual(hoja['A5'].value, 'Producto')
        self.assertEqual(hoja['C6'].value, 1050.0)
        self.assertEqual(hoja['C8'].value, '=SUBTOTAL(109,C6:C7)')
        # Un pedido con varios productos aparece en varias filas: no se suma.
        self.assertIsNone(hoja['D8'].value)
        self.assertEqual(len(hoja._charts), 1)
        self.assertIn('Análisis', libro.sheetnames)

    def test_markdown_a_lineas(self):
        lineas = _lineas_markdown('### Título\n- uno **dos** [link](x)\n2) tres')
        self.assertEqual(lineas, [('titulo', 'Título'), ('vineta', '- uno **dos** link'), ('vineta', '2) tres')])
